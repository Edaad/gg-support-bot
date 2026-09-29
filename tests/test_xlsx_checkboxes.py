"""Excel cell checkboxes added to openpyxl output."""

from __future__ import annotations

import io
import re
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from api.xlsx_checkboxes import CHECKBOX_NUMBER_FORMAT, apply_cell_checkboxes


def _xlsx(with_checkbox: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Category"
    ws["B1"] = 5
    ws["B1"].number_format = "$#,##0.00"
    if with_checkbox:
        for row in (2, 3):
            cell = ws.cell(row, 3, False)
            cell.number_format = CHECKBOX_NUMBER_FORMAT
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ApplyCellCheckboxesTestCase(unittest.TestCase):
    def test_marks_only_checkbox_cells(self) -> None:
        out = apply_cell_checkboxes(_xlsx(True))
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            styles = z.read("xl/styles.xml").decode()
            sheet = z.read("xl/worksheets/sheet1.xml").decode()
            types = z.read("[Content_Types].xml").decode()
            rels = z.read("xl/_rels/workbook.xml.rels").decode()
            bag = z.read("xl/featurePropertyBag/featurePropertyBag.xml").decode()
        xfs = re.search(r"<cellXfs.*?</cellXfs>", styles, re.DOTALL).group(0)
        xf_list = re.findall(r"<xf\b.*?(?:/>|</xf>)", xfs, re.DOTALL)
        marked = [i for i, xf in enumerate(xf_list) if "xfpb:xfComplement" in xf]
        self.assertEqual(len(marked), 1)
        self.assertIn('numFmtId="0"', xf_list[marked[0]])
        # Checkbox cells use the marked format; the currency cell does not.
        self.assertEqual(
            re.findall(r'<c r="C\d" s="(\d+)" t="b">', sheet), [str(marked[0])] * 2
        )
        self.assertNotIn(f's="{marked[0]}"', re.search(r'<c r="B1"[^>]*>', sheet)[0])
        self.assertIn("featurepropertybag+xml", types)
        self.assertEqual(rels.count("relationships/FeaturePropertyBag"), 1)
        self.assertIn('<bag type="Checkbox"/>', bag)
        # Still a readable workbook with boolean values.
        ws = load_workbook(io.BytesIO(out)).active
        self.assertIs(ws["C2"].value, False)

    def test_no_checkbox_cells_is_unchanged(self) -> None:
        raw = _xlsx(False)
        self.assertEqual(apply_cell_checkboxes(raw), raw)

    def test_idempotent(self) -> None:
        once = apply_cell_checkboxes(_xlsx(True))
        self.assertEqual(apply_cell_checkboxes(once), once)


if __name__ == "__main__":
    unittest.main()
