"""MCP Google Drive exports for bonus records and expenses."""

from __future__ import annotations

import io
import json
import os
import unittest
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch

import httpx
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import load_workbook

import mcp_server.db as mcp_db
from db.models import BonusRecord, BonusType, Club, Expense
from mcp_server import google_drive
from mcp_server.tools import exports
from tests.support.sqlite_db import make_sqlite_session_factory

TABLES = ["clubs", "bonus_types", "player_details", "bonus_records", "expenses"]
DRIVE_ENV = {
    "GOOGLE_OAUTH_CLIENT_ID": "client-id",
    "GOOGLE_OAUTH_CLIENT_SECRET": "client-secret",
    "GOOGLE_OAUTH_REFRESH_TOKEN": "refresh-token",
    "GDRIVE_EXPORT_FOLDER_ID": "folder-123",
}


class FakeDrive:
    """Records Drive calls and answers like the real API."""

    def __init__(self, *, token_status: int = 200, token_error: str | None = None):
        self.requests: list[httpx.Request] = []
        self.token_status = token_status
        self.token_error = token_error
        self.uploads: list[dict] = []

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        url = str(request.url)
        if url.startswith(google_drive.TOKEN_URL):
            if self.token_status >= 400:
                return httpx.Response(
                    self.token_status, json={"error": self.token_error or "bad"}
                )
            return httpx.Response(
                200, json={"access_token": "access-1", "expires_in": 3600}
            )
        if url.startswith(google_drive.UPLOAD_URL):
            body = request.content
            boundary = request.headers["content-type"].split("boundary=")[1]
            parts = body.split(f"--{boundary}".encode())
            metadata = json.loads(parts[1].split(b"\r\n\r\n", 1)[1].strip())
            media_head, media = parts[2].split(b"\r\n\r\n", 1)
            self.uploads.append(
                {
                    "metadata": metadata,
                    "media_type": media_head.decode().split("Content-Type: ")[1],
                    "media": media.rstrip(b"\r\n"),
                    "auth": request.headers["authorization"],
                }
            )
            return httpx.Response(
                200,
                json={
                    "id": "file-1",
                    "name": metadata["name"],
                    "webViewLink": "https://docs.google.com/spreadsheets/d/file-1/edit",
                },
            )
        if "/permissions" in url:
            return httpx.Response(200, json={"id": "anyoneWithLink"})
        return httpx.Response(404, json={"error": {"message": "unexpected"}})

    def client(self) -> httpx.Client:
        return httpx.Client(transport=httpx.MockTransport(self.handler))


class ExportTestBase(unittest.TestCase):
    def setUp(self) -> None:
        self.factory = make_sqlite_session_factory(TABLES)
        mcp_db.session_factory = self.factory
        self.addCleanup(setattr, mcp_db, "session_factory", None)
        env = patch.dict(os.environ, DRIVE_ENV, clear=False)
        env.start()
        self.addCleanup(env.stop)
        google_drive.reset_token_cache()
        self.addCleanup(google_drive.reset_token_cache)
        self.drive = FakeDrive()
        factory_patch = patch.object(google_drive, "client_factory", self.drive.client)
        factory_patch.start()
        self.addCleanup(factory_patch.stop)
        with self.factory() as s:
            s.add_all(
                [
                    Club(id=1, name="ClubGTO", telegram_user_id=1),
                    Club(id=2, name="Round Table", telegram_user_id=2),
                ]
            )
            s.add(BonusType(id=1, name="Rakeback", is_active=True, sort_order=0))
            s.add_all(
                [
                    BonusRecord(
                        player_username="GTO / 8190-5287 / Pirate",
                        group_title="GTO / 8190-5287 / Pirate",
                        gg_player_id="8190-5287",
                        amount=Decimal("25.50"),
                        bonus_type_id=1,
                        club_id=1,
                        issued_at=datetime(2026, 10, 6, 16, 0),
                    ),
                    BonusRecord(
                        player_username="GTO / 1111-2222 / Other",
                        amount=Decimal("10.00"),
                        custom_description="Goodwill, line one\nline two",
                        club_id=1,
                        issued_at=datetime(2026, 10, 7, 16, 0),
                    ),
                    BonusRecord(
                        player_username="RT / 3333-4444 / Table",
                        amount=Decimal("99.00"),
                        bonus_type_id=1,
                        club_id=2,
                        issued_at=datetime(2026, 10, 7, 16, 0),
                    ),
                    BonusRecord(
                        player_username="GTO / 5555-6666 / Late",
                        amount=Decimal("5.00"),
                        bonus_type_id=1,
                        club_id=1,
                        issued_at=datetime(2026, 10, 20, 16, 0),
                    ),
                ]
            )
            s.add_all(
                [
                    Expense(
                        amount=Decimal("120.00"),
                        expense_type="Software",
                        description="Hosting",
                        club_id=1,
                        expense_date=date(2026, 10, 1),
                        pending=True,
                    ),
                    Expense(
                        amount=Decimal("30.00"),
                        expense_type="Payout fee",
                        club_id=1,
                        expense_date=date(2026, 10, 2),
                        pending=False,
                    ),
                    Expense(
                        amount=Decimal("75.00"),
                        expense_type="Software",
                        club_id=2,
                        expense_date=date(2026, 10, 3),
                        pending=True,
                    ),
                ]
            )
            s.commit()


class BonusExportTests(ExportTestBase):
    def test_exports_filtered_bonuses_as_shared_sheet(self) -> None:
        result = exports.export_bonus_records(
            from_="2026-10-05", to="2026-10-11", club="clubgto"
        )
        item = result["item"]
        self.assertEqual(
            item["url"], "https://docs.google.com/spreadsheets/d/file-1/edit"
        )
        self.assertEqual(item["rows"], 2)
        self.assertEqual(item["total_amount"], 35.5)
        self.assertEqual(item["shared"], "anyone with the link (viewer)")
        self.assertFalse(result["pii_masked"])

        upload = self.drive.uploads[0]
        self.assertEqual(upload["metadata"]["mimeType"], google_drive.SHEET_MIME)
        self.assertEqual(upload["metadata"]["parents"], ["folder-123"])
        self.assertIn("ClubGTO", upload["metadata"]["name"])
        self.assertEqual(upload["media_type"], google_drive.CSV_MIME)
        self.assertEqual(upload["auth"], "Bearer access-1")
        csv_text = upload["media"].decode()
        self.assertIn("GTO / 8190-5287 / Pirate", csv_text)
        self.assertNotIn("RT / 3333-4444", csv_text)
        self.assertNotIn("Late", csv_text)

        permission = [r for r in self.drive.requests if "/permissions" in str(r.url)]
        self.assertEqual(len(permission), 1)
        self.assertEqual(
            json.loads(permission[0].content), {"type": "anyone", "role": "reader"}
        )

    def test_bonus_type_and_other_filters(self) -> None:
        typed = exports.export_bonus_records(
            from_="2026-10-05", to="2026-10-11", bonus_type_id=1
        )
        self.assertEqual(typed["item"]["rows"], 2)
        other = exports.export_bonus_records(
            from_="2026-10-05", to="2026-10-11", other=True
        )
        self.assertEqual(other["item"]["rows"], 1)
        self.assertEqual(other["item"]["total_amount"], 10.0)

    def test_bonus_validation(self) -> None:
        with self.assertRaises(ToolError):
            exports.export_bonus_records(from_="2026-10-11", to="2026-10-05")
        with self.assertRaises(ToolError):
            exports.export_bonus_records(
                from_="2026-10-05", to="2026-10-11", bonus_type_id=1, other=True
            )
        with self.assertRaises(ToolError):
            exports.export_bonus_records(from_="2026-10-05", to="2026-10-11", club="x")
        self.assertEqual(self.drive.uploads, [])

    def test_empty_export_notes_header_only(self) -> None:
        result = exports.export_bonus_records(from_="2025-01-01", to="2025-01-02")
        self.assertEqual(result["item"]["rows"], 0)
        self.assertTrue(any("header row" in n for n in result["notes"]))

    def test_row_cap(self) -> None:
        with patch.object(exports, "MAX_EXPORT_ROWS", 1):
            with self.assertRaises(ToolError):
                exports.export_bonus_records(from_="2026-10-05", to="2026-10-11")
        self.assertEqual(self.drive.uploads, [])


class ExpenseExportTests(ExportTestBase):
    def test_exports_filtered_expenses_as_xlsx_sheet(self) -> None:
        result = exports.export_expenses(club="clubgto", pending=True)
        item = result["item"]
        self.assertEqual(item["rows"], 1)
        self.assertEqual(item["total_amount"], 120.0)
        upload = self.drive.uploads[0]
        self.assertEqual(upload["media_type"], google_drive.XLSX_MIME)
        self.assertIn("pending", upload["metadata"]["name"])
        sheet = load_workbook(io.BytesIO(upload["media"])).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Amount")
        self.assertEqual(rows[1][1], "Software")
        self.assertEqual(len(rows), 2)

    def test_date_and_search_filters(self) -> None:
        result = exports.export_expenses(from_="2026-10-02", to="2026-10-03")
        self.assertEqual(result["item"]["rows"], 2)
        searched = exports.export_expenses(q="fee")
        self.assertEqual(searched["item"]["rows"], 1)

    def test_inverted_range(self) -> None:
        with self.assertRaises(ToolError):
            exports.export_expenses(from_="2026-10-03", to="2026-10-01")


class DriveClientTests(ExportTestBase):
    def test_token_is_cached_between_exports(self) -> None:
        exports.export_expenses()
        exports.export_expenses()
        token_calls = [
            r
            for r in self.drive.requests
            if str(r.url).startswith(google_drive.TOKEN_URL)
        ]
        self.assertEqual(len(token_calls), 1)
        self.assertIn(b"grant_type=refresh_token", token_calls[0].content)

    def test_missing_configuration(self) -> None:
        with patch.dict(os.environ, {"GOOGLE_OAUTH_REFRESH_TOKEN": ""}):
            with self.assertRaises(ToolError) as ctx:
                exports.export_expenses()
        self.assertIn("GOOGLE_OAUTH_REFRESH_TOKEN", str(ctx.exception))

    def test_revoked_refresh_token(self) -> None:
        self.drive.token_status = 400
        self.drive.token_error = "invalid_grant"
        with self.assertRaises(ToolError) as ctx:
            exports.export_expenses()
        self.assertIn("google_drive_auth.py", str(ctx.exception))

    def test_no_folder_uploads_to_root(self) -> None:
        with patch.dict(os.environ, {"GDRIVE_EXPORT_FOLDER_ID": ""}):
            exports.export_expenses()
        self.assertNotIn("parents", self.drive.uploads[0]["metadata"])


if __name__ == "__main__":
    unittest.main()
